import os
import json
import numpy as np
from tqdm import tqdm
import tensorflow as tf
from openpyxl import Workbook, load_workbook
from Utils.Decoders import generate_decoder
from Utils.Encoders import ImageEncoder, CaptionEncoder
from Utils.Fusion import generate_fusion
from Utils.EvaluationMetrics import (calculate_bleu_score, calculate_meteor_score,
                                     calculate_rouge_score, calculate_cider_score)
from Utils.Vectorization import generate_text_vectorization
from Utils.FeatureExtractors import load_batch_features, generate_features_filename
from Utils.WordEmbeddings import build_embedding_matrix

MAX_GRAD_NORM = 5.0


# The cross entropy loss function used to correlate images with captions
def multimodal_loss(labels, loss):
    mask = (labels != 0)
    mask = tf.squeeze(tf.cast(mask, loss.dtype))
    loss = loss * mask
    # Handle nan condition
    loss = tf.where(tf.math.is_nan(loss), tf.zeros_like(loss), loss)
    return tf.reduce_sum(loss) / tf.maximum(tf.reduce_sum(mask), 1.0)


# This method generates the results for the given predictions and reference captions
def evaluate_predictions(predictions=None, references=None):
    # Accumulate per-image scores
    bleu_1, bleu_2, bleu_3, bleu_4 = 0, 0, 0, 0
    meteor, rouge, cider = 0, 0, 0
    skipped_samples = 0

    for image, prediction in tqdm(predictions.items(), desc="CALCULATING EVALUATION SCORES"):
        # Fetch the ground truth references for the current image
        if image not in references:
            skipped_samples += 1
            continue
        ground_truth = references[image]

        # BLEU scores
        b_1, b_2, b_3, b_4 = calculate_bleu_score(references=ground_truth, predictions=prediction)
        bleu_1 += b_1
        bleu_2 += b_2
        bleu_3 += b_3
        bleu_4 += b_4

        # METEOR score
        meteor += calculate_meteor_score(references=ground_truth, predictions=prediction)

        # ROUGE-L score
        try:
            rouge += calculate_rouge_score(references=ground_truth, predictions=prediction)
        except ImportError:
            pass

        # CIDEr score
        cider += calculate_cider_score(references=ground_truth, predictions=prediction)

    # Calculate the average scores
    valid_samples = len(predictions) - skipped_samples
    if valid_samples > 0:
        bleu_1 /= valid_samples
        bleu_2 /= valid_samples
        bleu_3 /= valid_samples
        bleu_4 /= valid_samples
        meteor /= valid_samples
        rouge /= valid_samples
        cider /= valid_samples
    else:
        bleu_1, bleu_2, bleu_3, bleu_4 = 0, 0, 0, 0
        meteor, rouge, cider = 0, 0, 0

    # Return the evaluation results
    result = (f"\nBLEU 1:: {bleu_1}"
              f"\nBLEU 2:: {bleu_2}"
              f"\nBLEU 3:: {bleu_3}"
              f"\nBLEU 4:: {bleu_4}"
              f"\nMETEOR:: {meteor}"
              f"\nROUGE-L:: {rouge}"
              f"\nCIDEr:: {cider}")
    return result, (bleu_1, bleu_2, bleu_3, bleu_4, meteor, rouge, cider)


# This class holds all the model specific details with respect to the captioning pipeline
class ImageCaptioningModel(tf.keras.Model):
    # The loss function
    loss_object = tf.keras.losses.SparseCategoricalCrossentropy(from_logits=True,
                                                                reduction=tf.keras.losses.Reduction.NONE)

    def __init__(self, image_encoder=None, caption_encoder=None, decoder=None, fusion=None,
                 dataset=None, batch_size=None):
        super(ImageCaptioningModel, self).__init__()
        self.image_encoder = image_encoder
        self.caption_encoder = caption_encoder
        self.decoder = decoder
        self.fusion = fusion
        self.dataset = dataset
        self.batch_size = batch_size
        print(">>> Image Captioning Pipeline Initialized Successfully...")

    # Generates the caption for the given image file address
    def call(self, image_names=None, validation=False, beam_size=1):
        generated_captions = {}
        for image_name in image_names:
            image_features = load_batch_features([image_name])
            image_encoding = self.image_encoder(image_features, training=False)
            generated_captions[image_name] = [
                self.generate_caption_beam(image_encoding=image_encoding, beam_size=beam_size)
            ]
        return generated_captions

    # This method abstracts the training logic (renamed to avoid shadowing tf.keras.Model.train)
    def train_model(self, paths=None, start_epoch=0, epochs=None, save=False):
        # Load the previous states if start epoch is other than 0
        if start_epoch > 0:
            self.load_states(epoch=(start_epoch - 1), batch=0)

        # Track the optimized masked sequence loss for each batch.
        loss_list = []
        best_bleu_4 = 0
        # Start the epoch iteration
        for epoch in range(start_epoch, epochs):
            # Create batch loss list to save all the batch loss audit trails.
            batch_loss_list = []

            for (batch_number, (image_batch, captions_batch, caption_embeddings)) in tqdm(
                    tf.data.Dataset.enumerate(self.dataset.keras_dataset),
                    desc="BATCH TRAINING >>>", ncols=100):
                # Load the current batch image features
                batch_features = load_batch_features(image_batch.numpy())

                with tf.GradientTape() as tape:
                    # Encode the image features into the multimodal embeddings
                    visual_encoding = self.image_encoder(batch_features, training=True)
                    # Embed all teacher-forcing input tokens (positions 0 to caption_size-2) in one call
                    raw_input_embeddings = self.caption_encoder.encoding_layer(caption_embeddings[:, :-1])
                    # Fuse token inputs with the image, then keep the image as decoder memory.
                    if self.fusion is not None:
                        all_input_embeddings = self.fusion(visual_encoding, raw_input_embeddings, training=True)
                    else:
                        all_input_embeddings = raw_input_embeddings
                    # Apply dropout to input embeddings before passing to decoder
                    all_input_embeddings = self.caption_encoder.dropout_layer(all_input_embeddings, training=True)
                    # Single vectorized decoder forward pass — eliminates the growing hidden-state loop
                    all_predictions = self.decoder(all_input_embeddings, visual_encoding, training=True)
                    # all_predictions: (batch, caption_size-1, vocab_size)
                    loss = self.loss_object(caption_embeddings[:, 1:], all_predictions)
                    normalized_batch_loss = multimodal_loss(caption_embeddings[:, 1:], loss)
                    # Calculate all the trainable parameters of the pipeline
                    trainable_variables = (self.decoder.trainable_variables
                                           + self.image_encoder.trainable_variables
                                           + self.caption_encoder.trainable_variables)
                    # Include fusion trainable variables if fusion is active
                    if self.fusion is not None:
                        trainable_variables += self.fusion.trainable_variables
                    # Calculate gradients and clip to prevent exploding gradients
                    gradients = tape.gradient(normalized_batch_loss, trainable_variables)
                    gradients, _ = tf.clip_by_global_norm(gradients, MAX_GRAD_NORM)
                    self.optimizer.apply_gradients(grads_and_vars=zip(gradients, trainable_variables))
                    # Append the loss to the batch loss list
                    batch_loss_list.append(normalized_batch_loss.numpy())

            # ----------------------------------------------------------------------------------------------------------
            # VALIDATION TESTING (greedy decoding for speed during training)
            # ----------------------------------------------------------------------------------------------------------
            model_predictions = self.validation(paths=paths, image_names=self.dataset.validation_image_names,
                                                batch_size=128, beam_size=1)
            result, scores = evaluate_predictions(predictions=model_predictions,
                                                  references=self.dataset.validation_image_captions)
            # Save the validation results to the file
            self.save_generated_captions(epoch_number=epoch, captions=model_predictions)
            self.save_validation_results(epoch_number=epoch, scores=scores)

            if save:
                self.save_states(epoch, 0)
                if scores[3] > best_bleu_4:
                    best_bleu_4 = scores[3]
                    print(">>> BLEU-4 improved to", best_bleu_4, "- model saved.")
                else:
                    print(">>> BLEU-4 did not improve. Current:", scores[3], "Best:", best_bleu_4, "- model saved.")
            # ----------------------------------------------------------------------------------------------------------

            # Append the optimized masked sequence loss to the list for visualization.
            loss_list.append(tf.reduce_sum(batch_loss_list).numpy())
            # Save the current epoch loss list to the file
            self.save_losses(epoch_number=epoch, epoch_loss=tf.reduce_sum(batch_loss_list).numpy(),
                             batch_losses=batch_loss_list)
            print("EPOCH NUMBER: ", epoch, " TOTAL LOSS: ", loss_list[epoch - start_epoch])

    # This method abstracts the fine-tuning logic
    def finetune(self, paths=None, epochs=None, save=False, existing_scores=None):
        # Track the optimized masked sequence loss for each batch.
        loss_list = []
        best_bleu_4 = existing_scores[3] if existing_scores else 0
        # Start the epoch iteration
        for epoch in range(epochs):
            # Create batch loss list to save all the batch loss audit trails.
            batch_loss_list = []

            for (batch_number, (image_batch, captions_batch, caption_embeddings)) in tqdm(
                    tf.data.Dataset.enumerate(self.dataset.keras_dataset),
                    desc="BATCH TRAINING >>>", ncols=100):
                # Load the current batch image features
                batch_features = load_batch_features(image_batch.numpy())

                with tf.GradientTape() as tape:
                    # Encode the image features into the multimodal embeddings
                    visual_encoding = self.image_encoder(batch_features, training=False)
                    # Embed all teacher-forcing input tokens in one call
                    raw_input_embeddings = self.caption_encoder.encoding_layer(caption_embeddings[:, :-1])
                    # Fusion is frozen during fine-tuning (training=False) — only decoder weights are updated
                    if self.fusion is not None:
                        all_input_embeddings = self.fusion(visual_encoding, raw_input_embeddings, training=False)
                    else:
                        all_input_embeddings = raw_input_embeddings
                    # Apply dropout to input embeddings before passing to decoder
                    all_input_embeddings = self.caption_encoder.dropout_layer(all_input_embeddings, training=True)
                    # Single vectorized decoder forward pass
                    all_predictions = self.decoder(all_input_embeddings, visual_encoding, training=True)
                    loss = self.loss_object(caption_embeddings[:, 1:], all_predictions)
                    normalized_batch_loss = multimodal_loss(caption_embeddings[:, 1:], loss)
                    # During fine-tuning, freeze all layers except the decoder (fusion frozen per DHARA)
                    decoder_trainables = self.decoder.trainable_variables
                    # Calculate gradients and clip to prevent exploding gradients
                    gradients = tape.gradient(normalized_batch_loss, decoder_trainables)
                    gradients, _ = tf.clip_by_global_norm(gradients, MAX_GRAD_NORM)
                    self.optimizer.apply_gradients(grads_and_vars=zip(gradients, decoder_trainables))
                    # Append the loss to the batch loss list
                    batch_loss_list.append(normalized_batch_loss.numpy())

            # ----------------------------------------------------------------------------------------------------------
            # TESTING
            # ----------------------------------------------------------------------------------------------------------
            model_predictions = self.validation(paths=paths, image_names=self.dataset.testing_image_names,
                                                batch_size=128, beam_size=5)
            result, scores = evaluate_predictions(predictions=model_predictions,
                                                  references=self.dataset.testing_image_captions)

            print("\nFinetuning Epoch ", str(epoch), " ", result)

            if scores[3] > best_bleu_4 and save:
                best_bleu_4 = scores[3]
                self.save_states(epoch, 0, finetuning=True)
                self.save_generated_captions(epoch_number=epoch, captions=model_predictions, finetuning=True)
                self.save_validation_results(epoch_number=epoch, scores=scores, finetuning=True)
            # ----------------------------------------------------------------------------------------------------------

            # Append the optimized masked sequence loss to the list for visualization.
            loss_list.append(tf.reduce_sum(batch_loss_list).numpy())
            print("EPOCH NUMBER: ", epoch, " TOTAL FINETUNING LOSS: ", loss_list[epoch])

    # This method validates the model over the validation split
    def validation(self, paths=None, image_names=None, batch_size=None, beam_size=None):
        # Directory of the image features
        directory = paths.features_abs_path()
        feature_names = [generate_features_filename(directory_path=directory, file_name=image) for image in image_names]

        # Create a dataset version to access in batches
        image_names_dataset = tf.data.TextLineDataset.from_tensor_slices(feature_names)
        image_names_dataset = image_names_dataset.batch(batch_size=batch_size).prefetch(
            buffer_size=tf.data.experimental.AUTOTUNE)

        validation_generated_captions = {}

        for image_batch in tqdm(image_names_dataset, desc="EVALUATING SPLIT >>>", ncols=100):
            batch_features = load_batch_features(image_batch.numpy())
            batch_image_encoding = self.image_encoder(batch_features, training=False)

            for index, image in enumerate(image_batch):
                image_key = os.path.basename(image.numpy().decode()).replace(".npy", ".jpg")
                image_encoding = batch_image_encoding[index:index + 1]
                caption = self.generate_caption_beam(image_encoding=image_encoding, beam_size=beam_size)
                validation_generated_captions[image_key] = [caption]

        return validation_generated_captions

    def generate_caption_beam(self, image_encoding=None, beam_size=5, length_penalty=0.7):
        start_id = int(self.word_to_index("startsen").numpy())
        end_id = int(self.word_to_index("endsen").numpy())
        beams = [([start_id], 0.0, self.decoder.reset_states(1), False)]

        for token_number in range(1, self.caption_size):
            candidates = []
            for token_ids, score, hidden, ended in beams:
                if ended:
                    candidates.append((token_ids, score, hidden, ended))
                    continue

                input_token = tf.constant([token_ids[-1]], dtype=tf.int64)
                input_embeddings = self.caption_encoder(input_token, training=False)
                if self.fusion is not None:
                    decoder_input = self.fusion(
                        image_encoding, tf.expand_dims(input_embeddings, axis=1), training=False)[:, -1, :]
                else:
                    decoder_input = input_embeddings

                predictions, next_hidden = self.decoder(decoder_input, image_encoding, hidden=hidden, training=False)
                predictions = self.suppress_invalid_tokens(predictions=predictions, token_number=token_number)
                log_probs = tf.nn.log_softmax(predictions[0]).numpy()
                top_indices = np.argsort(log_probs)[-beam_size:][::-1]

                for index in top_indices:
                    index = int(index)
                    next_ids = token_ids + [index]
                    candidates.append((next_ids, score + float(log_probs[index]), next_hidden, index == end_id))

            beams = sorted(candidates, key=lambda item: self.beam_score(item[1], len(item[0]), length_penalty),
                           reverse=True)[:beam_size]
            if all(ended for _, _, _, ended in beams):
                break

        best_tokens, _, _, _ = max(beams, key=lambda item: self.beam_score(item[1], len(item[0]), length_penalty))
        words = self.decode_token_ids(best_tokens, start_id=start_id, end_id=end_id)
        return " ".join(words)

    def beam_score(self, log_probability=None, token_count=None, length_penalty=0.7):
        generated_length = max(token_count - 1, 1)
        return log_probability / (generated_length ** length_penalty)

    def decode_token_ids(self, token_ids=None, start_id=None, end_id=None):
        words = []
        for token_id in token_ids:
            if token_id == start_id:
                continue
            if token_id == end_id:
                break
            word = self.index_to_word(tf.constant([token_id])).numpy()[0].decode()
            if word and word != "[UNK]":
                words.append(word)
        return words

    def suppress_invalid_tokens(self, predictions=None, token_number=None, min_caption_tokens=2):
        logits = predictions.numpy().copy()
        blocked_tokens = ["", "[UNK]", "startsen"]
        if token_number <= min_caption_tokens:
            blocked_tokens.append("endsen")
        blocked_indices = self.word_to_index(tf.constant(blocked_tokens)).numpy()
        blocked_indices = [int(index) for index in blocked_indices if int(index) < logits.shape[-1]]
        if blocked_indices:
            logits[:, blocked_indices] = -np.inf
        return tf.convert_to_tensor(logits, dtype=predictions.dtype)

    def evaluate_split_loss(self, paths=None, image_captions=None, batch_size=128):
        directory = paths.features_abs_path()
        feature_paths, token_sequences = [], []

        for image_name, captions in image_captions.items():
            feature_path = generate_features_filename(directory_path=directory, file_name=image_name)
            for caption in captions:
                feature_paths.append(feature_path)
                token_sequences.append(self.caption_to_token_ids(caption))

        losses = []
        for start in tqdm(range(0, len(feature_paths), batch_size), desc="CALCULATING TEACHER-FORCED LOSS",
                          ncols=100):
            end = start + batch_size
            batch_features = load_batch_features(feature_paths[start:end])
            batch_tokens = tf.constant(token_sequences[start:end], dtype=tf.int64)
            visual_encoding = self.image_encoder(batch_features, training=False)
            raw_input_embeddings = self.caption_encoder.encoding_layer(batch_tokens[:, :-1])
            if self.fusion is not None:
                all_input_embeddings = self.fusion(visual_encoding, raw_input_embeddings, training=False)
            else:
                all_input_embeddings = raw_input_embeddings
            all_predictions = self.decoder(all_input_embeddings, visual_encoding, training=False)
            loss = self.loss_object(batch_tokens[:, 1:], all_predictions)
            losses.append(multimodal_loss(batch_tokens[:, 1:], loss).numpy())

        return float(np.mean(losses)) if losses else 0.0

    def caption_to_token_ids(self, caption=None):
        tokens = ["startsen"] + str(caption).split() + ["endsen"]
        vocabulary = set(self.vocabulary)
        tokens = [token for token in tokens if token in vocabulary]
        token_ids = self.word_to_index(tf.constant(tokens)).numpy().astype(np.int64).tolist()
        token_ids = token_ids[:self.caption_size]
        if len(token_ids) < self.caption_size:
            token_ids.extend([0] * (self.caption_size - len(token_ids)))
        return token_ids

    # This method saves the states of the model subcomponents
    def save_states(self, epoch=None, batch=None, finetuning=False):
        # Generate filenames to store model states
        suffix = f"Batch_{batch}_FineTuned_Epoch_{epoch}" if finetuning else f"Batch_{batch}_Epoch_{epoch}"

        image_encoder_filepath = os.path.join(self.paths.image_encoder_saving_path(), suffix)
        caption_encoder_filepath = os.path.join(self.paths.caption_encoder_saving_path(), suffix)
        decoder_filepath = os.path.join(self.paths.decoder_saving_path(), suffix)
        fusion_filepath = os.path.join(self.paths.fusion_saving_path(), suffix)

        # Ensure all state directories exist before saving
        for filepath in [image_encoder_filepath, caption_encoder_filepath, decoder_filepath, fusion_filepath]:
            os.makedirs(os.path.dirname(filepath), exist_ok=True)

        # Save the model states now
        self.image_encoder.save_weights(filepath=image_encoder_filepath, overwrite=True)
        self.caption_encoder.save_weights(filepath=caption_encoder_filepath, overwrite=True)
        self.decoder.save_weights(filepath=decoder_filepath, overwrite=True)
        if self.fusion is not None:
            self.fusion.save_weights(filepath=fusion_filepath, overwrite=True)

    # This method checks whether the state exists or not
    def state_exists(self, epoch=None, batch=0):
        image_encoder_filepath = os.path.join(self.paths.image_encoder_saving_path(), f"Batch_{batch}_Epoch_{epoch}")
        try:
            self.image_encoder.load_weights(image_encoder_filepath)
        except (OSError, ValueError, tf.errors.NotFoundError):
            return False
        return True

    # This method loads the states back
    def load_states(self, epoch=None, batch=None):
        # Generate filenames to load the model states
        suffix = f"Batch_{batch}_Epoch_{epoch}"
        image_encoder_filepath = os.path.join(self.paths.image_encoder_saving_path(), suffix)
        caption_encoder_filepath = os.path.join(self.paths.caption_encoder_saving_path(), suffix)
        decoder_filepath = os.path.join(self.paths.decoder_saving_path(), suffix)
        fusion_filepath = os.path.join(self.paths.fusion_saving_path(), suffix)

        # Load the model states now
        self.image_encoder.load_weights(image_encoder_filepath)
        self.caption_encoder.load_weights(caption_encoder_filepath)
        self.decoder.load_weights(decoder_filepath)
        if self.fusion is not None:
            self.fusion.load_weights(fusion_filepath)

    # This method saves the training logs to an Excel file for visualizations
    def save_losses(self, epoch_number=None, epoch_loss=None, batch_losses=None):
        training_log_filepath = self.paths.training_logs_saving_path()
        if os.path.exists(training_log_filepath):
            wb = load_workbook(training_log_filepath)
            ws = wb.active
            # Add more batch columns to the header if this epoch has more batches than previous ones
            current_batch_cols = ws.max_column - 2  # subtract Epoch and Epoch Loss columns
            if len(batch_losses) > current_batch_cols:
                for i in range(current_batch_cols + 1, len(batch_losses) + 1):
                    ws.cell(row=1, column=i + 2, value=f"Batch {i}")
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Training Logs"
            header = ["Epoch", "Epoch Loss"] + [f"Batch {i}" for i in range(1, len(batch_losses) + 1)]
            ws.append(header)
        row = [epoch_number, epoch_loss] + list(batch_losses)
        ws.append(row)
        wb.save(training_log_filepath)

    # This method saves the validation results to an Excel file for visualizations
    def save_validation_results(self, epoch_number=None, scores=None, finetuning=False):
        prediction_filepath = self.paths.predictions_saving_path()
        if os.path.exists(prediction_filepath):
            wb = load_workbook(prediction_filepath)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Prediction Results"
            ws.append(["Epoch Number", "BLEU-1", "BLEU-2", "BLEU-3", "BLEU-4", "METEOR", "ROUGE", "CIDEr"])
        bleu_1, bleu_2, bleu_3, bleu_4, meteor, rouge, cider = scores
        ws.append([epoch_number, bleu_1, bleu_2, bleu_3, bleu_4, meteor, rouge, cider])
        wb.save(prediction_filepath)

    # This method saves the generated captions to the external text file
    def save_generated_captions(self, captions=None, epoch_number=None, finetuning=False):
        filepath = self.paths.generated_predictions_saving_path()
        with open(file=filepath, mode="a", encoding='utf-8') as caption_logs:
            mark = "\n----------------------------------------------------------------------------------------------\n"
            epoch_number_text = " Finetuning Epoch Number " if finetuning else " Epoch Number "
            epoch_ = mark + epoch_number_text + str(epoch_number) + mark
            caption_logs.write(epoch_)
            json.dump(captions, caption_logs, indent=4, ensure_ascii=False)
            caption_logs.write(mark)

    def save_test_results(self, epoch=None, batch=0, scores=None, teacher_forced_loss=None):
        output_dir = self.paths.test_results_dir()
        os.makedirs(output_dir, exist_ok=True)
        filepath = os.path.join(output_dir, f"TestResults_Batch_{batch}_Epoch_{epoch}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Results"
        ws.append(["Epoch", "Batch", "Teacher Forced Loss", "BLEU-1", "BLEU-2", "BLEU-3",
                   "BLEU-4", "METEOR", "ROUGE", "CIDEr"])
        bleu_1, bleu_2, bleu_3, bleu_4, meteor, rouge, cider = scores
        ws.append([epoch, batch, teacher_forced_loss, bleu_1, bleu_2, bleu_3, bleu_4, meteor, rouge, cider])
        wb.save(filepath)

    def save_test_captions(self, epoch=None, batch=0, captions=None):
        output_dir = self.paths.test_results_dir()
        os.makedirs(output_dir, exist_ok=True)
        filepath = os.path.join(output_dir, f"TestCaptions_Batch_{batch}_Epoch_{epoch}.txt")
        with open(filepath, mode="w", encoding="utf-8") as caption_logs:
            json.dump(captions, caption_logs, indent=4, ensure_ascii=False)


# This method generates an instance of the image captioning model for the CLIP ViT + Transformer pipeline
def initialize_captioning_model(paths=None, batch_size=None, image_feature_shape=None, embedding_size=None,
                                dataset=None, word_embedding_type=None, decoder_type=None,
                                fusion_type=None, **kwargs):
    # Here, we set the vocabulary length only to the words known in the threshold training captions
    text_vectorization = generate_text_vectorization(vocabulary=dataset.vocabulary,
                                                     max_caption_length=dataset.max_caption_length)
    vocabulary_length = len(text_vectorization.get_vocabulary())

    # Build pretrained embedding matrix if a pretrained type is selected
    pretrained_weights = None
    if word_embedding_type is not None:
        pretrained_weights, pretrained_dim = build_embedding_matrix(
            vocabulary=dataset.vocabulary,
            embedding_type=word_embedding_type,
            project_root=paths.project_path())
        # Override embedding_size with the pretrained dimension
        embedding_size = pretrained_dim
        print(f">>> Embedding size overridden to {embedding_size} (from pretrained {word_embedding_type})")

    # Generate the image encoder
    image_encoder = ImageEncoder(batch_size=batch_size, input_shape=image_feature_shape,
                                 encoding_size=embedding_size)

    # Generate the caption encoder (with or without pretrained weights)
    caption_encoder = CaptionEncoder(batch_size=batch_size, vocabulary_length=vocabulary_length,
                                     encoding_size=embedding_size, pretrained_weights=pretrained_weights)

    # Generate the TransformerDecoder
    decoder = generate_decoder(batch_size=batch_size, input_shape=embedding_size, output_shape=vocabulary_length)

    # Generate the multimodal fusion block if configured (uses embedding_size as d_model)
    fusion = generate_fusion(fusion_type=fusion_type, d_model=embedding_size)

    # Now, every subcomponent of the pipeline have been created, generate the model now
    model = ImageCaptioningModel(image_encoder=image_encoder, caption_encoder=caption_encoder,
                                 decoder=decoder, fusion=fusion, dataset=dataset, batch_size=batch_size)

    # Set the vocabulary, vocabulary length and caption size, batch size, paths
    model.vocabulary = text_vectorization.get_vocabulary()
    model.vocabulary_length = vocabulary_length
    model.caption_size = dataset.max_caption_length
    model.paths = paths

    # Add the word to index and index to word mappings and text vectorization layers to the model
    model.word_to_index = tf.keras.layers.StringLookup(mask_token="", vocabulary=dataset.vocabulary)
    model.index_to_word = tf.keras.layers.StringLookup(mask_token="", vocabulary=dataset.vocabulary, invert=True)

    # Return the generated model
    return model
